"""Tests for the multi-format ingest layer (.docx / .md).

PDF + TXT extraction is already exercised indirectly through the gate tests.
Here we focus on the two new paths and the ``_gather_books`` tie-breaker
(PDF wins over DOCX wins over MD wins over TXT for the same stem).
"""

from __future__ import annotations

from pathlib import Path


def _write_md(path: Path) -> None:
    path.write_text(
        "# Глава 1\n\n"
        "**Стейкхолдер** — это _заинтересованное лицо_ "
        "(см. [статью](https://example.com/stakeholders)).\n\n"
        "```python\n"
        "# code fences must be stripped\n"
        "def foo(): pass\n"
        "```\n\n"
        "## Подраздел\n\n"
        "Вторая мысль про стейкхолдеров и их влияние на организацию.\n",
        encoding="utf-8",
    )


def _write_docx(path: Path) -> None:
    """Build a minimal .docx with two paragraphs and a table using
    python-docx — same dependency the extractor uses."""
    from docx import Document

    doc = Document()
    doc.add_paragraph("Стейкхолдер — это физическое или юридическое лицо.")
    doc.add_paragraph("Заинтересованные стороны включают акционеров, поставщиков, клиентов.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Тип стейкхолдера"
    table.rows[0].cells[1].text = "Акционеры"
    doc.save(str(path))


# ---------- markdown ----------


def test_extract_md_strips_fences_and_markup(tmp_path) -> None:
    from src.rag.ingest import _extract_md

    md = tmp_path / "chapter.md"
    _write_md(md)
    pages = _extract_md(md)
    assert pages, "expected at least one page"
    joined = "\n".join(p for _, p in pages)
    # Fenced code must not survive.
    assert "def foo" not in joined
    assert "```" not in joined
    # Inline **bold** / _italics_ get unwrapped but the content stays.
    assert "Стейкхолдер" in joined
    assert "заинтересованное лицо" in joined
    # Link text is preserved, URL is dropped.
    assert "статью" in joined
    assert "example.com" not in joined
    # Headings lose the '#' prefix but keep the heading text.
    assert "Глава 1" in joined
    assert "#" not in joined


def test_extract_md_empty_file_yields_nothing(tmp_path) -> None:
    from src.rag.ingest import _extract_md

    md = tmp_path / "empty.md"
    md.write_text("", encoding="utf-8")
    assert _extract_md(md) == []


# ---------- docx ----------


def test_extract_docx_paragraphs_and_tables(tmp_path) -> None:
    from src.rag.ingest import _extract_docx

    path = tmp_path / "notes.docx"
    _write_docx(path)
    pages = _extract_docx(path)
    assert pages, "expected at least one page"
    joined = "\n".join(p for _, p in pages)
    assert "Стейкхолдер" in joined
    assert "акционеров" in joined
    # Table contents must flow through — we had a bug where only
    # ``doc.paragraphs`` was read and tables were silently dropped.
    assert "Тип стейкхолдера" in joined
    assert "Акционеры" in joined


def test_extract_docx_missing_paragraphs_returns_empty(tmp_path) -> None:
    from docx import Document

    from src.rag.ingest import _extract_docx

    path = tmp_path / "empty.docx"
    Document().save(str(path))
    assert _extract_docx(path) == []


# ---------- gather_books tie-breaking ----------


def test_gather_books_prefers_pdf_over_txt(tmp_path) -> None:
    from src.rag.ingest import _gather_books

    (tmp_path / "ch1.pdf").write_bytes(b"%PDF-1.4")
    (tmp_path / "ch1.txt").write_text("plain", encoding="utf-8")
    picked = _gather_books(tmp_path)
    assert [p.name for p in picked] == ["ch1.pdf"]


def test_gather_books_prefers_docx_over_md_and_txt(tmp_path) -> None:
    from src.rag.ingest import _gather_books

    (tmp_path / "ch1.md").write_text("hi", encoding="utf-8")
    (tmp_path / "ch1.docx").write_bytes(b"PK")  # payload content irrelevant
    (tmp_path / "ch1.txt").write_text("hi", encoding="utf-8")
    picked = _gather_books(tmp_path)
    assert [p.name for p in picked] == ["ch1.docx"]


def test_gather_books_keeps_distinct_stems(tmp_path) -> None:
    from src.rag.ingest import _gather_books

    (tmp_path / "intro.pdf").write_bytes(b"%PDF")
    (tmp_path / "appendix.md").write_text("x", encoding="utf-8")
    (tmp_path / "notes.docx").write_bytes(b"PK")
    picked = sorted(p.name for p in _gather_books(tmp_path))
    assert picked == ["appendix.md", "intro.pdf", "notes.docx"]


def test_gather_books_ignores_unknown_extensions(tmp_path) -> None:
    from src.rag.ingest import _gather_books

    (tmp_path / "image.png").write_bytes(b"\x89PNG")
    (tmp_path / "data.csv").write_text("a,b", encoding="utf-8")
    (tmp_path / "ch.txt").write_text("hi", encoding="utf-8")
    picked = [p.name for p in _gather_books(tmp_path)]
    assert picked == ["ch.txt"]


# ---------- dispatcher ----------


def test_extract_book_dispatch(tmp_path) -> None:
    from src.rag.ingest import _extract_book

    md = tmp_path / "a.md"
    md.write_text("Hello **world**\n\nSecond paragraph.\n", encoding="utf-8")
    txt = tmp_path / "b.txt"
    txt.write_text("Plain text page.", encoding="utf-8")
    docx = tmp_path / "c.docx"
    _write_docx(docx)

    md_pages = _extract_book(md)
    txt_pages = _extract_book(txt)
    docx_pages = _extract_book(docx)

    assert md_pages and "world" in md_pages[0][1]
    assert txt_pages and "Plain text page." in txt_pages[0][1]
    assert docx_pages and "Стейкхолдер" in docx_pages[0][1]
